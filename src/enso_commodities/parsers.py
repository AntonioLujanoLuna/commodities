from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SEASON_TO_MONTH = {
    "DJF": 1,
    "JFM": 2,
    "FMA": 3,
    "MAM": 4,
    "AMJ": 5,
    "MJJ": 6,
    "JJA": 7,
    "JAS": 8,
    "ASO": 9,
    "SON": 10,
    "OND": 11,
    "NDJ": 12,
}
PERIOD_PATTERN = re.compile(r"^(\d{4})M(\d{2})$")


@dataclass(frozen=True)
class PinkSheetDiagnostics:
    source_updated: str | None
    first_period: str
    last_period: str
    months: int
    commodities: int
    mismatch_rows: int
    non_positive_values: int


def parse_noaa_index(path: Path, *, name: str) -> pd.DataFrame:
    if name not in {"roni", "oni"}:
        raise ValueError("name must be 'roni' or 'oni'")
    raw = pd.read_csv(path, sep=r"\s+")
    expected = {"SEAS", "YR", "ANOM"}
    if not expected.issubset(raw.columns):
        raise ValueError(f"{path}: expected NOAA columns {sorted(expected)}")
    if not raw["SEAS"].isin(SEASON_TO_MONTH).all():
        invalid = sorted(set(raw.loc[~raw["SEAS"].isin(SEASON_TO_MONTH), "SEAS"]))
        raise ValueError(f"{path}: invalid seasons: {invalid}")

    result = pd.DataFrame(
        {
            "date": pd.to_datetime(
                {
                    "year": raw["YR"].astype(int),
                    "month": raw["SEAS"].map(SEASON_TO_MONTH).astype(int),
                    "day": 1,
                }
            ),
            name: pd.to_numeric(raw["ANOM"], errors="raise"),
            f"{name}_season": raw["SEAS"].astype(str),
        }
    )
    if result["date"].duplicated().any():
        raise ValueError(f"{path}: duplicate NOAA month")
    if not result[name].between(-5.0, 5.0).all():
        raise ValueError(f"{path}: implausible {name.upper()} anomaly outside [-5, 5]")
    if len(result) < 800:
        raise ValueError(f"{path}: unexpectedly short NOAA history ({len(result)} rows)")
    return result.sort_values("date", ignore_index=True)


def parse_pink_sheet(path: Path) -> tuple[pd.DataFrame, PinkSheetDiagnostics]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Monthly Prices" not in workbook.sheetnames:
        raise ValueError(f"{path}: 'Monthly Prices' worksheet not found")
    rows = list(workbook["Monthly Prices"].iter_rows(values_only=True))

    first_data_index = next(
        (
            index
            for index, row in enumerate(rows)
            if isinstance(row[0], str) and PERIOD_PATTERN.fullmatch(row[0].strip())
        ),
        None,
    )
    if first_data_index is None or first_data_index < 2:
        raise ValueError(f"{path}: monthly price data block not found")

    names = rows[first_data_index - 2]
    units = rows[first_data_index - 1]
    commodity_columns = [
        index for index, value in enumerate(names) if index > 0 and isinstance(value, str)
    ]
    if len(commodity_columns) < 50:
        raise ValueError(f"{path}: only {len(commodity_columns)} named commodity columns")

    records: list[dict[str, object]] = []
    periods: list[str] = []
    for row in rows[first_data_index:]:
        period = row[0]
        if not isinstance(period, str) or not PERIOD_PATTERN.fullmatch(period.strip()):
            continue
        period = period.strip()
        match = PERIOD_PATTERN.fullmatch(period)
        assert match is not None
        year, month = map(int, match.groups())
        if not 1 <= month <= 12:
            raise ValueError(f"{path}: invalid month in {period}")
        periods.append(period)
        timestamp = pd.Timestamp(year=year, month=month, day=1)
        for column in commodity_columns:
            source_value = pd.to_numeric(
                row[column] if column < len(row) else None, errors="coerce"
            )
            non_positive = pd.notna(source_value) and float(source_value) <= 0
            records.append(
                {
                    "date": timestamp,
                    "period": period,
                    "commodity": str(names[column]).strip(),
                    "unit": str(units[column]).strip() if units[column] is not None else None,
                    "source_value": source_value,
                    "value": pd.NA if non_positive else source_value,
                    "quality_flag": "source_non_positive" if non_positive else None,
                    "source_column": column + 1,
                }
            )

    result = pd.DataFrame.from_records(records)
    if result.empty:
        raise ValueError(f"{path}: no monthly price observations parsed")
    if result.duplicated(["date", "commodity"]).any():
        raise ValueError(f"{path}: duplicate date/commodity observations")
    if result["date"].nunique() < 700:
        raise ValueError(f"{path}: unexpectedly short price history")
    result["value"] = pd.to_numeric(result["value"], errors="coerce")

    updated = next(
        (
            str(row[0]).strip()
            for row in rows[:10]
            if row and isinstance(row[0], str) and row[0].startswith("Updated on ")
        ),
        None,
    )
    mismatch_rows = 0
    if "Mismatch Details" in workbook.sheetnames:
        mismatch_rows = max(workbook["Mismatch Details"].max_row - 4, 0)
    diagnostics = PinkSheetDiagnostics(
        source_updated=updated,
        first_period=min(periods),
        last_period=max(periods),
        months=len(set(periods)),
        commodities=len(commodity_columns),
        mismatch_rows=mismatch_rows,
        non_positive_values=int(result["quality_flag"].notna().sum()),
    )
    return result.sort_values(["date", "commodity"], ignore_index=True), diagnostics


def parse_pink_sheet_total_index(path: Path) -> pd.DataFrame:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Monthly Indices" not in workbook.sheetnames:
        raise ValueError(f"{path}: 'Monthly Indices' worksheet not found")
    rows = list(workbook["Monthly Indices"].iter_rows(values_only=True))
    first_data_index = next(
        (
            index
            for index, row in enumerate(rows)
            if isinstance(row[0], str) and PERIOD_PATTERN.fullmatch(row[0].strip())
        ),
        None,
    )
    if first_data_index is None:
        raise ValueError(f"{path}: monthly index data block not found")
    header_rows = rows[max(0, first_data_index - 5) : first_data_index]
    total_index_column = next(
        (
            column
            for row in header_rows
            for column, value in enumerate(row)
            if value == "Total Index"
        ),
        None,
    )
    if total_index_column is None:
        raise ValueError(f"{path}: 'Total Index' column not found")

    records: list[dict[str, object]] = []
    for row in rows[first_data_index:]:
        period = row[0]
        if not isinstance(period, str) or not PERIOD_PATTERN.fullmatch(period.strip()):
            continue
        match = PERIOD_PATTERN.fullmatch(period.strip())
        assert match is not None
        year, month = map(int, match.groups())
        value = pd.to_numeric(row[total_index_column], errors="coerce")
        records.append(
            {
                "date": pd.Timestamp(year=year, month=month, day=1),
                "period": period.strip(),
                "world_bank_total_index": value,
            }
        )
    result = pd.DataFrame.from_records(records)
    if len(result) < 700:
        raise ValueError(f"{path}: unexpectedly short World Bank total-index history")
    if result["date"].duplicated().any():
        raise ValueError(f"{path}: duplicate World Bank total-index month")
    if result["world_bank_total_index"].isna().any():
        raise ValueError(f"{path}: missing World Bank total-index value")
    if not result["world_bank_total_index"].gt(0).all():
        raise ValueError(f"{path}: non-positive World Bank total-index value")
    return result.sort_values("date", ignore_index=True)
